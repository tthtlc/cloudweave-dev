#!/usr/bin/env python3
"""Generate the libcloud_nutanix system-overview deck.

Run with the venv that has python-pptx:
    /tmp/pptxenv/bin/python presentation/generate_deck.py

Every fact in this deck is source-verified; see ARCHITECTURE.md at the repo root
and the per-subsystem ARCHITECTURE.md files. No secret VALUES appear anywhere —
credentials are referred to by variable name only.
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- palette ---------------------------------------------------------------
INK = RGBColor(0x0F, 0x17, 0x2A)      # near-black slate
MUTED = RGBColor(0x64, 0x74, 0x8B)    # secondary text
PAPER = RGBColor(0xF8, 0xFA, 0xFC)    # slide background
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1E, 0x29, 0x3B)     # dark panels / title slide
BLUE = RGBColor(0x25, 0x63, 0xEB)     # primary accent
CYAN = RGBColor(0x06, 0x94, 0xA2)     # secondary accent
GREEN = RGBColor(0x05, 0x96, 0x69)    # good / correct
AMBER = RGBColor(0xD9, 0x77, 0x06)    # caution
RED = RGBColor(0xDC, 0x26, 0x26)      # critical
LINE = RGBColor(0xCB, 0xD5, 0xE1)     # hairlines
CHIP = RGBColor(0xE2, 0xE8, 0xF0)     # chip fill

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)
BODY_W = W - 2 * MARGIN

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# --- primitives ------------------------------------------------------------
def _txbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _para(tf, text, size, color, bold=False, space_before=0, space_after=6,
          align=PP_ALIGN.LEFT, font="Calibri", first=False, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return p


def _rect(slide, x, y, w, h, fill, line_col=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
          adj=0.08):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line_col is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_col
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = adj
        except (IndexError, KeyError):
            pass
    return s


def _blank(bg=PAPER):
    s = prs.slides.add_slide(BLANK)
    bgfill = s.background.fill
    bgfill.solid()
    bgfill.fore_color.rgb = bg
    return s


def slide_header(slide, title, kicker=None):
    """Standard content-slide header: kicker + title + accent rule."""
    y = Inches(0.46)
    if kicker:
        tf = _txbox(slide, MARGIN, y, BODY_W, Inches(0.24))
        _para(tf, kicker.upper(), 11.5, BLUE, bold=True, space_after=0, first=True)
        y += Inches(0.30)
    tf = _txbox(slide, MARGIN, y, BODY_W, Inches(0.55))
    _para(tf, title, 29, INK, bold=True, space_after=0, first=True)
    rule = _rect(slide, MARGIN, y + Inches(0.58), Inches(1.05), Pt(3.5), BLUE,
                 shape=MSO_SHAPE.RECTANGLE)
    return rule.top + rule.height + Inches(0.30)


def footer(slide, text, n):
    tf = _txbox(slide, MARGIN, H - Inches(0.5), BODY_W - Inches(0.6), Inches(0.25))
    _para(tf, text, 9.5, MUTED, space_after=0, first=True)
    tf2 = _txbox(slide, W - MARGIN - Inches(0.6), H - Inches(0.5), Inches(0.6), Inches(0.25))
    _para(tf2, str(n), 9.5, MUTED, align=PP_ALIGN.RIGHT, space_after=0, first=True)


# --- slide builders --------------------------------------------------------
def title_slide(title, subtitle, meta):
    s = _blank(NAVY)
    _rect(s, Inches(0), Inches(0), W, Pt(6), BLUE, shape=MSO_SHAPE.RECTANGLE)
    tf = _txbox(s, MARGIN, Inches(2.25), Inches(10.4), Inches(1.4))
    _para(tf, title, 46, WHITE, bold=True, space_after=10, first=True, line=1.02)
    tf2 = _txbox(s, MARGIN, Inches(3.85), Inches(9.6), Inches(0.9))
    _para(tf2, subtitle, 19, RGBColor(0x94, 0xA3, 0xB8), space_after=0, first=True,
          line=1.25)
    _rect(s, MARGIN, Inches(5.05), Inches(0.9), Pt(3), CYAN, shape=MSO_SHAPE.RECTANGLE)
    tf3 = _txbox(s, MARGIN, Inches(5.45), Inches(10.0), Inches(0.6))
    _para(tf3, meta, 12.5, RGBColor(0x64, 0x74, 0x8B), space_after=0, first=True)
    return s


def section_slide(num, title, blurb, n):
    s = _blank(NAVY)
    _rect(s, Inches(0), Inches(0), Pt(7), H, BLUE, shape=MSO_SHAPE.RECTANGLE)
    tf = _txbox(s, Inches(1.15), Inches(2.7), Inches(1.6), Inches(0.7))
    _para(tf, f"{num:02d}", 54, RGBColor(0x33, 0x41, 0x55), bold=True, space_after=0,
          first=True)
    tf2 = _txbox(s, Inches(2.6), Inches(2.78), Inches(9.5), Inches(0.8))
    _para(tf2, title, 34, WHITE, bold=True, space_after=8, first=True)
    tf3 = _txbox(s, Inches(2.6), Inches(3.72), Inches(9.0), Inches(0.9))
    _para(tf3, blurb, 15, RGBColor(0x94, 0xA3, 0xB8), space_after=0, first=True,
          line=1.3)
    tf4 = _txbox(s, W - MARGIN - Inches(0.6), H - Inches(0.5), Inches(0.6), Inches(0.25))
    _para(tf4, str(n), 9.5, RGBColor(0x47, 0x55, 0x69), align=PP_ALIGN.RIGHT,
          space_after=0, first=True)
    return s


def bullets_slide(title, kicker, items, n, foot="", lead=None):
    """items: list of (text, level) or (text, level, color)."""
    s = _blank()
    y = slide_header(s, title, kicker)
    if lead:
        tf = _txbox(s, MARGIN, y, BODY_W, Inches(0.5))
        _para(tf, lead, 15.5, MUTED, space_after=0, first=True, line=1.28)
        y += Inches(0.30) + Inches(0.26) * (1 + len(lead) // 115)
    tf = _txbox(s, MARGIN, y, BODY_W, H - y - Inches(0.75))
    first = True
    for item in items:
        text, level = item[0], item[1]
        color = item[2] if len(item) > 2 else (INK if level == 0 else MUTED)
        size = 16.5 if level == 0 else 14
        bullet = "▪  " if level == 0 else "–  "
        p = _para(tf, bullet + text, size, color, bold=(level == 0),
                  space_before=(9 if level == 0 and not first else 2),
                  space_after=3, first=first, line=1.22)
        p.level = level
        first = False
    footer(s, foot, n)
    return s


def table_slide(title, kicker, headers, rows, n, foot="", widths=None,
                lead=None, accent_col=None):
    s = _blank()
    y = slide_header(s, title, kicker)
    if lead:
        tf = _txbox(s, MARGIN, y, BODY_W, Inches(0.4))
        _para(tf, lead, 14.5, MUTED, space_after=0, first=True, line=1.25)
        y += Inches(0.46)
    nrows, ncols = len(rows) + 1, len(headers)
    avail_h = H - y - Inches(0.75)
    row_h = min(Inches(0.42), avail_h / nrows)
    tbl_h = row_h * nrows
    gt = s.shapes.add_table(nrows, ncols, MARGIN, y, BODY_W, tbl_h).table
    gt.first_row = True
    if widths:
        total = sum(widths)
        for i, wgt in enumerate(widths):
            gt.columns[i].width = Emu(int(BODY_W * wgt / total))
    for r in range(nrows):
        gt.rows[r].height = Emu(int(row_h))
    for c, htext in enumerate(headers):
        cell = gt.cell(0, c)
        cell.text = htext
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.09)
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = WHITE
                run.font.name = "Calibri"
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = gt.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else RGBColor(0xF1, 0xF5, 0xF9)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.09)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11.5)
                    run.font.color.rgb = INK
                    run.font.name = "Calibri"
                    if accent_col is not None and c == accent_col:
                        run.font.bold = True
                        run.font.color.rgb = BLUE
                    if c == 0 and accent_col is None:
                        run.font.bold = True
    footer(s, foot, n)
    return s


def _chip(slide, x, y, w, h, label, sub, fill, text_col=WHITE, sub_col=None,
          label_sz=13, sub_sz=9.5):
    _rect(slide, x, y, w, h, fill)
    tf = _txbox(slide, x + Inches(0.08), y + Inches(0.10), w - Inches(0.16),
                h - Inches(0.18))
    _para(tf, label, label_sz, text_col, bold=True, space_after=1, first=True,
          align=PP_ALIGN.CENTER)
    if sub:
        _para(tf, sub, sub_sz, sub_col or RGBColor(0xCB, 0xD5, 0xE1), space_after=0,
              align=PP_ALIGN.CENTER)


def topology_slide(n):
    s = _blank()
    y = slide_header(s, "Deployment topology", "One Docker bridge network")
    tf = _txbox(s, MARGIN, y, BODY_W, Inches(0.45))
    _para(tf, "All containers share the external bridge network libcloud_net and address each other by "
              "service DNS name. The portal's nginx is the only intended public surface.",
          14.5, MUTED, space_after=0, first=True, line=1.25)
    y += Inches(0.62)

    # browser
    _chip(s, MARGIN, y + Inches(1.05), Inches(1.55), Inches(0.85), "Browser",
          "end user", RGBColor(0x47, 0x55, 0x69))

    # portal (front door)
    px = MARGIN + Inches(1.95)
    _rect(s, px, y + Inches(0.30), Inches(2.35), Inches(2.35), BLUE)
    tf = _txbox(s, px + Inches(0.10), y + Inches(0.45), Inches(2.15), Inches(2.0))
    _para(tf, "portal", 15, WHITE, bold=True, space_after=2, first=True,
          align=PP_ALIGN.CENTER)
    _para(tf, "nginx + React SPA", 10, RGBColor(0xBF, 0xDB, 0xFE), space_after=8,
          align=PP_ALIGN.CENTER)
    _para(tf, ":3000", 19, WHITE, bold=True, space_after=8, align=PP_ALIGN.CENTER)
    _para(tf, "/api/  → identity", 10, RGBColor(0xBF, 0xDB, 0xFE), space_after=1,
          align=PP_ALIGN.CENTER)
    _para(tf, "/dex/  → dex", 10, RGBColor(0xBF, 0xDB, 0xFE), space_after=1,
          align=PP_ALIGN.CENTER)
    _para(tf, "/      → SPA", 10, RGBColor(0xBF, 0xDB, 0xFE), space_after=0,
          align=PP_ALIGN.CENTER)

    # loopback panel
    lx = px + Inches(2.75)
    panel_w = Inches(6.05)
    _rect(s, lx - Inches(0.12), y + Inches(0.02), panel_w + Inches(0.24),
          Inches(3.55), RGBColor(0xEF, 0xF3, 0xF8), LINE)
    tf = _txbox(s, lx, y + Inches(0.12), panel_w, Inches(0.24))
    _para(tf, "BOUND TO 127.0.0.1  —  NOT REACHABLE OFF-HOST", 10, MUTED, bold=True,
          space_after=0, first=True)

    cw, ch, gapx, gapy = Inches(1.85), Inches(0.78), Inches(0.25), Inches(0.22)
    row1 = y + Inches(0.46)
    _chip(s, lx, row1, cw, ch, "identity-service", ":8766  sessions + authz", NAVY)
    _chip(s, lx + cw + gapx, row1, cw, ch, "dex", ":5556  OIDC issuer", NAVY)
    _chip(s, lx + 2 * (cw + gapx), row1, cw, ch, "lldap", ":3890  directory", NAVY)

    row2 = row1 + ch + gapy
    _chip(s, lx, row2, cw, ch, "libcloud-rest-api", ":8765  cloud API", NAVY)
    _chip(s, lx + cw + gapx, row2, cw, ch, "openfga", ":8080  decisions", NAVY)
    _chip(s, lx + 2 * (cw + gapx), row2, cw, ch, "vault", ":8200  secrets", NAVY)

    row3 = row2 + ch + gapy
    _chip(s, lx, row3, cw, ch, "openfga-postgres", ":5432  tuple store", NAVY)
    _chip(s, lx + cw + gapx, row3, cw, ch, "prism ×4", ":4010-4013  specs", NAVY)
    _chip(s, lx + 2 * (cw + gapx), row3, cw, ch, "AWS / Nutanix", "real backends",
          RGBColor(0x47, 0x55, 0x69))

    # all-interfaces strip
    sy = y + Inches(3.75)
    tf = _txbox(s, MARGIN, sy, BODY_W, Inches(0.24))
    _para(tf, "ALSO PUBLISHED ON ALL INTERFACES (0.0.0.0)", 10, RED, bold=True,
          space_after=0, first=True)
    _chip(s, MARGIN, sy + Inches(0.28), Inches(2.9), Inches(0.62), "portal  :3000",
          "intended public surface", BLUE)
    _chip(s, MARGIN + Inches(3.10), sy + Inches(0.28), Inches(2.9), Inches(0.62),
          "openfga-visualizer  :5050", "dev server, superadmin-gated", AMBER)
    _chip(s, MARGIN + Inches(6.20), sy + Inches(0.28), Inches(3.6), Inches(0.62),
          "nutanix emulator  :9440-9443", "no authentication at all", RED)

    footer(s, "Source: docker-compose.yml across all subprojects; server/Dockerfile:30-57", n)
    return s


def flow_slide(title, kicker, steps, n, foot=""):
    """steps: list of (actor, detail). Rendered as a numbered vertical sequence."""
    s = _blank()
    y = slide_header(s, title, kicker)
    row_h = min(Inches(0.60), (H - y - Inches(0.8)) / len(steps))
    for i, (actor, detail) in enumerate(steps):
        ry = y + i * row_h
        _rect(s, MARGIN, ry + Inches(0.03), Inches(0.34), Inches(0.34), BLUE,
              shape=MSO_SHAPE.OVAL)
        tf = _txbox(s, MARGIN, ry + Inches(0.07), Inches(0.34), Inches(0.28))
        _para(tf, str(i + 1), 11, WHITE, bold=True, align=PP_ALIGN.CENTER,
              space_after=0, first=True)
        tf2 = _txbox(s, MARGIN + Inches(0.5), ry, Inches(3.15), row_h)
        _para(tf2, actor, 13.5, INK, bold=True, space_after=0, first=True)
        tf3 = _txbox(s, MARGIN + Inches(3.75), ry, BODY_W - Inches(3.75), row_h)
        _para(tf3, detail, 12.5, MUTED, space_after=0, first=True, line=1.15,
              font="Consolas")
    footer(s, foot, n)
    return s


def findings_slide(title, kicker, findings, n, foot=""):
    """findings: list of (severity, color, headline, detail)."""
    s = _blank()
    y = slide_header(s, title, kicker)
    row_h = (H - y - Inches(0.8)) / len(findings)
    for sev, col, head, detail in findings:
        _rect(s, MARGIN, y + Inches(0.04), Pt(4), row_h - Inches(0.16), col,
              shape=MSO_SHAPE.RECTANGLE)
        tf = _txbox(s, MARGIN + Inches(0.22), y + Inches(0.02), Inches(1.15),
                    Inches(0.26))
        _para(tf, sev, 10.5, col, bold=True, space_after=0, first=True)
        tf2 = _txbox(s, MARGIN + Inches(1.45), y, BODY_W - Inches(1.45), Inches(0.3))
        _para(tf2, head, 15, INK, bold=True, space_after=0, first=True)
        tf3 = _txbox(s, MARGIN + Inches(1.45), y + Inches(0.30), BODY_W - Inches(1.45),
                     row_h - Inches(0.34))
        _para(tf3, detail, 12, MUTED, space_after=0, first=True, line=1.2)
        y += row_h
    footer(s, foot, n)
    return s


def closing_slide(title, lines, n):
    s = _blank(NAVY)
    _rect(s, Inches(0), Inches(0), W, Pt(6), BLUE, shape=MSO_SHAPE.RECTANGLE)
    tf = _txbox(s, MARGIN, Inches(1.5), Inches(11.5), Inches(0.8))
    _para(tf, title, 34, WHITE, bold=True, space_after=0, first=True)
    _rect(s, MARGIN, Inches(2.42), Inches(0.9), Pt(3), CYAN, shape=MSO_SHAPE.RECTANGLE)
    tf2 = _txbox(s, MARGIN, Inches(2.85), Inches(11.5), Inches(3.4))
    first = True
    for text, kind in lines:
        if kind == "h":
            _para(tf2, text, 15, CYAN, bold=True, space_before=12, space_after=4,
                  first=first)
        else:
            _para(tf2, text, 13.5, RGBColor(0x94, 0xA3, 0xB8), space_after=3,
                  first=first, line=1.25)
        first = False
    tf3 = _txbox(s, W - MARGIN - Inches(0.6), H - Inches(0.5), Inches(0.6), Inches(0.25))
    _para(tf3, str(n), 9.5, RGBColor(0x47, 0x55, 0x69), align=PP_ALIGN.RIGHT,
          space_after=0, first=True)
    return s


# ===========================================================================
# Deck content
# ===========================================================================
FOOT = "libcloud_nutanix — system architecture"
n = 0

# 1 ─ title
n += 1
title_slide(
    "Multi-Tenant Cloud\nProvisioning Platform",
    "Self-service VM provisioning on AWS and Nutanix, with federated identity,\n"
    "relationship-based authorization, and centrally held cloud credentials.",
    "System architecture walkthrough  ·  All facts verified against source",
)

# 2 ─ what it is
n += 1
bullets_slide(
    "What the system does", "Overview",
    [
        ("A user signs in through a browser and provisions VMs on AWS or Nutanix", 0),
        ("They never hold cloud credentials — the platform holds them in Vault and acts on their behalf", 1),
        ("Two tenants ship seeded: aws and nutanix", 0),
        ("Roles per tenant: owner, admin, viewer — plus a platform-wide superadmin", 1),
        ("Eight LLDAP users are created at bootstrap, including cloud-denied", 1),
        ("cloud-denied exists to prove denial works — authenticated, but authorized for nothing", 1),
        ("Each tenant now gets its own Vault AppRole identity, mapped in OpenFGA", 0),
        ("Five concerns are split across five services, deliberately", 0),
        ("No single component both decides permission and holds the credential", 1),
    ],
    n, FOOT,
    lead="A provisioning portal where authentication, authorization, and secrets are three separate systems.",
)

# 3 ─ separation of concerns
n += 1
table_slide(
    "Separation of concerns", "Overview",
    ["Concern", "Component", "Answers the question"],
    [
        ["Identity", "LLDAP + Dex", "Which human is this?"],
        ["Session", "identity-service", "Is this browser still that human?"],
        ["Authorization", "OpenFGA", "May this human do this to that?"],
        ["Secrets", "Vault", "What credential provisions this tenant?"],
        ["Cloud abstraction", "libcloud.rest + Apache libcloud", "Talk to AWS / Nutanix"],
    ],
    n, FOOT, widths=[0.20, 0.32, 0.48],
    lead="The split is the point — it is what keeps a compromised component from being a total compromise.",
)

# 4 ─ topology
n += 1
topology_slide(n)

# 5 ─ section: the three surfaces
n += 1
section_slide(1, "The three surfaces",
              "What a person can actually open in a browser: the portal, the "
              "authorization visualizer, and the Nutanix emulator.", n)

# 6 ─ portal
n += 1
bullets_slide(
    "Port 3000 — the Role Portal", "Surface 1 of 3",
    [
        ("React 18 SPA (Create React App), built to a static bundle, served by nginx:alpine", 0),
        ("Multi-stage build — no Node runtime in the shipped image", 1),
        ("nginx is the single front door, and this is the key architectural fact", 0),
        ("/api/  →  identity-service:8766      /dex/  →  dex:5556      /  →  SPA fallback", 1),
        ("The browser never contacts Dex, the identity service, or the REST API directly", 1),
        ("This is WHY every other service binds to 127.0.0.1", 1),
        ("Role-specific dashboards: Viewer, Admin, Owner, SuperAdmin", 0),
        ("Plus state pages: PendingApproval, DisabledAccount, Unauthorized, IdentityCollapse", 1),
        ("The session is an httpOnly cookie — the browser stores no tokens", 0),
        ("sessionStorage holds UX metadata only; every call uses credentials: \"include\"", 1),
        ("Client-side route guards are cosmetic; enforcement is server-side", 1),
    ],
    n, FOOT,
)

# 7 ─ visualizer
n += 1
bullets_slide(
    "Port 5050 — the OpenFGA Visualizer", "Surface 2 of 3",
    [
        ("A single-file Flask app that makes the authorization model legible", 0),
        ("Five tabs: Hierarchy, Permission Matrix, Check, Model, Model Graph", 1),
        ("D3 collapsible tree for relationships; vis-network force graph for the model itself", 1),
        ("Read-only against OpenFGA — it can read, check and expand, but never write", 0),
        ("Answers \"why can this user do this?\" by expanding the resolution path", 1),
        ("Well-guarded, despite being a side tool", 0),
        ("Its own Dex SSO login, full JWT verification, server-side sessions", 1),
        ("Hard-gated to the LLDAP superadmin only — no other user can sign in", 1),
        ("But: published on 0.0.0.0:5050 on a Werkzeug development server, over plain HTTP", 0),
        ("The superadmin gate limits exposure; a dev server on a public interface is still not production", 1),
    ],
    n, FOOT,
)

# 8 ─ emulator
n += 1
bullets_slide(
    "Ports 9440-9443 — the Nutanix Emulator", "Surface 3 of 3",
    [
        ("A two-tier stand-in for Prism Central v4, so the stack runs with no real cluster", 0),
        ("Tier 1: a stateful Node/Express shim on :9440, HTTPS with a self-signed cert", 1),
        ("Tier 2: Stoplight Prism on :4010, serving schema-valid examples for unhandled paths", 1),
        ("Four API minor versions run side by side", 0),
        ("v4.0 → :9440    v4.1 → :9441    v4.2 → :9442    v4.3 → :9443", 1),
        ("The v4.0 merged spec carries 487 paths, 2 206 schemas, 109 tags", 1),
        ("Fidelity gaps that change how you read test results", 0),
        ("Authentication is an explicit no-op — any credential is accepted", 1),
        ("No ETag / If-Match, so the driver's concurrency control is silently untested", 1),
        ("stop_node sends \"shutdown\", which the shim ignores — success is returned, nothing happens", 1),
        ("Tasks always succeed 400 ms after creation; there are no failure paths", 1),
    ],
    n, FOOT,
    lead="Useful for development. Treat every green test against it with the gaps below in mind.",
)

# 9 ─ section: identity
n += 1
section_slide(2, "Authentication",
              "Who the user is: LLDAP as the directory, Dex as the OIDC issuer, "
              "and three credentials that are not interchangeable.", n)

# 10 ─ dex/lldap
n += 1
table_slide(
    "LLDAP and Dex", "Identity",
    ["Property", "Value"],
    [
        ["Directory", "LLDAP — base DN dc=libcloud,dc=local, users in ou=people"],
        ["Issuer (iss)", "http://dex:5556/dex   — the in-container DNS name, deliberately"],
        ["Dex storage", "memory — all tokens are lost on restart"],
        ["Connector", "lldap, LDAP to lldap:3890, insecureNoSSL: true"],
        ["Claim mapping", "uid → sub,  mail → email,  cn → name"],
        ["Groups claim", "none — no groupSearch is configured"],
        ["Grants enabled", "authorization_code, refresh_token  (no password, no client_credentials)"],
        ["OAuth clients", "libcloud-portal (browser, PKCE S256)  ·  libcloud-rest (services, secret)"],
    ],
    n, FOOT, widths=[0.22, 0.78], accent_col=None,
    lead="The issuer is the container DNS name so every service can validate tokens and fetch JWKS on the internal network.",
)

# 11 ─ three credentials
n += 1
table_slide(
    "Three credentials, not interchangeable", "Identity",
    ["Credential", "Who holds it", "Why it exists"],
    [
        ["Session cookie\nlibcloud_portal_sid",
         "The browser",
         "HS256 JWT minted by the identity service. HttpOnly, SameSite=Lax, 8 h.\nThe Dex refresh token stays server-side and never reaches the browser."],
        ["Provisioner token\naud = libcloud-rest",
         "identity-service",
         "The portal user's token has the wrong audience for the REST API, so the\nservice logs into Dex as aws-admin / ntnx-admin and mints a second token."],
        ["Superadmin JWT\naud = libcloud-rest",
         "setup.sh, operators",
         "The bootstrap gate. Obtained by a real OIDC login, not a static bypass."],
    ],
    n, FOOT, widths=[0.24, 0.20, 0.56],
    lead="Misreading which credential is in play is the easiest way to misunderstand this system.",
)

# 12 ─ the trust boundary
n += 1
bullets_slide(
    "The consequence: identity-service is the trust boundary", "Identity",
    [
        ("The libcloud REST API sees the provisioner service account, not the end user", 0),
        ("It re-checks authorization — but against aws-admin / ntnx-admin, not the human", 1),
        ("The human's permissions are evaluated in the identity service, and nowhere else on that path", 0),
        ("If that check is wrong or skipped, nothing downstream catches it", 1),
        ("When the identity service calls OpenFGA, two different principals are in play", 0),
        ("The bearer token authenticates the CALLER (the provisioner)", 1),
        ("The tuple's user field names the SUBJECT being evaluated (the end user)", 1),
        ("This is correct OpenFGA usage — but it is easy to misread as a privilege confusion", 1),
    ],
    n, FOOT,
)

# 13 ─ section: authorization
n += 1
section_slide(3, "Authorization",
              "OpenFGA holds every permission decision as relationships, "
              "not as roles baked into code.", n)

# 14 ─ the model
n += 1
bullets_slide(
    "The authorization model", "OpenFGA",
    [
        ("OpenFGA v1.16.0, PostgreSQL-backed, 9 types on schema 1.1", 0),
        ("user · platform · tenant · libcloud_api · provider · resource_class · aws_region · nutanix_cluster · vault_user", 1),
        ("Bootstrap seeds 50 tuples — 41 structural wiring + 9 role grants", 1),
        ("A new vault_user type maps each tenant to its Vault AppRole identity", 0),
        ("tenant:<t> parent vault_user:libcloud-<t> — the tenant's Vault user", 1),
        ("superadmin is a control-plane role, deliberately", 0),
        ("It grants global_reader (read everything) and the right to assign tenant owners", 1),
        ("It does NOT grant can_provision anywhere — it cannot create resources", 1),
        ("Backend writes require an INTERSECTION — this is the tenant kill switch", 0),
        ("can_provision = (tenant_admin or tenant_owner) AND can_use from provider", 1),
        ("Revoke a tenant's provider access and provisioning stops instantly for everyone in it,", 1),
        ("with no per-user tuples touched", 1),
        ("Version pin v1.16.0 is load-bearing", 0),
        ("Earlier releases don't refetch JWKS on an unknown key id, so Dex's 6-hourly", 1),
        ("key rotation breaks every check until OpenFGA is restarted", 1),
    ],
    n, FOOT,
)

# 15 ─ enforcement points
n += 1
table_slide(
    "Three enforcement points", "OpenFGA",
    ["#", "Where", "What it checks", "For whom"],
    [
        ["1", "identity-service,\nper verb",
         "can_read / can_provision / can_update on\naws_region:aws or nutanix_cluster:nutanix",
         "The END USER\n(from the session cookie)"],
        ["2", "libcloud.rest,\nper route",
         "Scope from policies.json, then can_connect,\ncan_use, and can_provision / can_read",
         "The PROVISIONER\n(defence in depth)"],
        ["3", "Vault credential\nwrite",
         "can_manage_credentials on tenant:<t>",
         "The tenant OWNER"],
    ],
    n, FOOT, widths=[0.05, 0.19, 0.48, 0.28],
    lead="Point 1 is the one that reflects the human's permissions. Point 2 guards the service account.",
)

# 16 ─ policy table
n += 1
bullets_slide(
    "The REST API's policy table", "Authorization",
    [
        ("Authorization is not decorated onto handlers — handlers contain none at all", 0),
        ("A custom route class, AuthorizedAPIRoute, enforces before the handler runs", 1),
        ("Rules live in an external, hot-reloadable file: app/auth/policies.json", 0),
        ("Keyed \"METHOD /path/template\", carrying scopes, an authz scope, and a driver capability", 1),
        ("Reload without restart via POST /v1/admin/policies:reload", 1),
        ("It is fail-closed, and that is the right default", 0),
        ("A route with no policy entry returns 500 policy_unknown_operation", 1),
        ("Adding an endpoint without a policy makes it unreachable, not unprotected", 1),
    ],
    n, FOOT,
)

# 17 ─ section: flows
n += 1
section_slide(4, "Request flows",
              "What actually happens on the wire, from the login click "
              "to a running virtual machine.", n)

# 18 ─ login flow
n += 1
flow_slide(
    "Browser login", "Authorization code + PKCE",
    [
        ("Browser → portal", "GET :3000/login   → nginx SPA fallback"),
        ("Browser → identity", "GET /api/auth/begin?provider=lldap&redirect_uri=..."),
        ("identity-service", "mints state (32B) + PKCE verifier (48B), stores server-side, TTL 600 s"),
        ("Browser → Dex", "GET :3000/dex/auth ... &code_challenge=<S256>&connector_id=lldap"),
        ("Browser → Dex", "POST /dex/auth/lldap/login   → Dex binds to LLDAP, matches uid"),
        ("Dex → Browser", "302 to :3000/auth/callback?code=...&state=..."),
        ("Browser → identity", "POST /api/auth/exchange { provider, code, state, redirectUri }"),
        ("identity-service", "consume_state() — single-use pop. THE authoritative CSRF check"),
        ("identity → Dex", "POST /dex/token   grant=authorization_code + secret + code_verifier"),
        ("identity-service", "verify id_token: JWKS sig, iss, aud=libcloud-portal, exp"),
        ("identity → Browser", "Set-Cookie: libcloud_portal_sid=<HS256 JWT>; HttpOnly; SameSite=Lax"),
    ],
    n, "The PKCE verifier and the client secret both stay server-side. The browser handles neither.",
)

# 19 ─ provisioning flow
n += 1
flow_slide(
    "Provisioning a VM", "End to end",
    [
        ("Browser → identity", "POST /api/provision/aws { vmName }   (cookie)"),
        ("identity-service", "verify cookie JWT (HS256, exp)"),
        ("identity → OpenFGA", "POST /stores/<id>/check   user:<END USER>  can_provision  aws_region:aws"),
        ("", "bearer = provisioner token · subject = the human   → deny here ends it"),
        ("identity → Dex", "cached server-side LDAP login as aws-admin → aud=libcloud-rest"),
        ("identity → REST", "GET /v1/auth/me · POST /v1/connections:test"),
        ("identity → REST", "GET locations · sizes · images · nodes · subnets"),
        ("", "X-Provider-Connection: { provider, region, auth_binding }  — NO credentials"),
        ("libcloud.rest", "policy lookup → JWT verify → OpenFGA can_connect / can_use / can_provision"),
        ("REST → Vault", "resolve vault_user (OpenFGA) → AppRole material → AppRole login → read secret (60 m token, 30 s cache)"),
        ("REST → cloud", "Apache libcloud driver → AWS EC2 or Nutanix Prism Central"),
    ],
    n, "The client names auth_binding; the server resolves the tenant's Vault AppRole and fetches the value.",
)

# 20 ─ secrets
n += 1
bullets_slide(
    "Secrets: Vault", "Credentials",
    [
        ("Vault 1.15, file storage, KV v2 at secret/ + AppRole auth at approle/", 0),
        ("Per-tenant cloud secret: secret/data/libcloud/<tenant>", 1),
        ("aws → access key / secret        nutanix → Prism username / password", 1),
        ("Each tenant has its own Vault identity — an AppRole, not a shared token", 0),
        ("AppRole libcloud-<tenant>, policy libcloud-read-<tenant> (read only that tenant)", 1),
        ("The tenant → vault-user mapping lives in OpenFGA (vault_user type)", 1),
        ("The single global read token is replaced by a narrow orchestrator token", 0),
        ("It can read AppRole auth material only — never the cloud secrets", 1),
        ("Read path: resolve vault_user → AppRole login → 60 m tenant token → read", 1),
        ("Writing is still owner-gated via OpenFGA can_manage_credentials", 0),
        ("Client-supplied credentials are rejected with 403 unless explicitly enabled", 1),
        ("Development-grade posture: single-share unseal, root token on disk, TLS disabled", 0),
    ],
    n, FOOT,
)

# 21 ─ bootstrap
n += 1
bullets_slide(
    "The bootstrap chain of trust", "Operations",
    [
        ("setup.sh establishes trust in a strict order, each step gated on the last", 0),
        ("Steps 1-3: render Dex config → start LLDAP, create superadmin → start Postgres, OpenFGA, Dex, Vault", 1),
        ("Step 4 is the pivot — a REAL OIDC login as superadmin", 0),
        ("It produces SUPERADMIN_JWT, verified inside the identity-service container", 1),
        ("Everything privileged downstream refuses to run without it", 1),
        ("Steps 5-7: create tenant users · seed the OpenFGA model and tuples · bootstrap Vault", 0),
        ("Bootstrap is not a backdoor — it uses the same identity path a human would", 1),
        ("Companion scripts", 0),
        ("rebuild_all.sh rebuilds images online; setup.sh itself is build-free and works air-gapped", 1),
        ("docker_teardown.sh is host-wide and destroys volumes — unseal key and tuples are lost", 1),
    ],
    n, FOOT,
)

# 22 ─ section: security
n += 1
section_slide(5, "Security posture",
              "What the design gets right, and the findings that need "
              "attention before this is exposed.", n)

# 23 ─ done well
n += 1
bullets_slide(
    "What the design gets right", "Security",
    [
        ("Cloud credentials never reach the browser or the client", 0, GREEN),
        ("The client names a binding; the server resolves the value from Vault", 1),
        ("The Dex refresh token is held server-side and never sent to the browser", 0, GREEN),
        ("PKCE S256 with a server-held verifier, plus server-issued single-use state", 0, GREEN),
        ("Complete ID-token validation: JWKS signature, iss, aud, exp, required claims", 0, GREEN),
        ("The REST API's policy table is fail-closed", 0, GREEN),
        ("Authorization is enforced server-side on every request", 0, GREEN),
        ("The SPA's route guards are cosmetic, and the code says so", 1),
        ("New federated users land with no tuples — authenticated, authorized for nothing", 0, GREEN),
        ("Bootstrap privilege is gated on a real OIDC login, not a static bypass", 0, GREEN),
        ("Most services bind to loopback, with the portal as a deliberate single ingress", 0, GREEN),
    ],
    n, FOOT,
)

# 24 ─ findings 1
n += 1
findings_slide(
    "Findings — the composing three", "Security",
    [
        ("CRITICAL", RED, "Live secrets are committed to git — eleven tracked files",
         "Vault root token and unseal key, all eight LLDAP passwords, OAuth client secrets, the LDAP bind "
         "password and AWS keys. Two are easy to miss: a Markdown walkthrough and a captured stderr log "
         "with credentials pasted inline. Untracking is not enough — the values are in history."),
        ("CRITICAL", RED, "The session signing key is a placeholder bootstrap never generates",
         "SESSION_SECRET stays at \"change-me\". Unlike the Postgres password and Dex client secrets, no "
         "setup path generates it. The session cookie is an HS256 JWT signed with it — so it is forgeable. "
         "The same key also signs the identity-collapse pending token."),
        ("CRITICAL", RED, "Role checks trust the cookie's self-asserted role",
         "_require_role returns before consulting OpenFGA when the cookie already claims the required role "
         "or superadmin. A forged cookie claiming superadmin reaches user management, raw tuple CRUD, and "
         "the OpenFGA explorer with no authorization lookup at all."),
    ],
    n, "Individually critical. Together they are a path from \"read the repo\" to platform admin.",
)

# 25 ─ findings 2
n += 1
findings_slide(
    "Findings — availability and exposure", "Security",
    [
        ("HIGH", AMBER, "A one-shot discovery failure silently disables authorization",
         "The identity service latches its \"discovered\" flag BEFORE attempting OpenFGA discovery, and "
         "catches bare Exception. One transient failure leaves it permanently disabled — and the disabled "
         "path returns allow. There is no depends_on, so a cold start is a live race."),
        ("HIGH", AMBER, "OpenFGA authenticates but does not authorize its own API",
         "OIDC validates issuer, audience and signature — it does not scope what a caller may do. Any user "
         "who can obtain a libcloud-rest-audience token and reach openfga:8080 can write tuples and grant "
         "themselves any relation. The only control is network placement."),
        ("HIGH", AMBER, "No TLS anywhere in the stack",
         "Dex issues over http, the portal serves http, LDAP runs insecureNoSSL, Vault has tls_disable. "
         "The session cookie is therefore necessarily Secure=false. Session cookies, authorization codes "
         "and Vault tokens all cross the wire in clear text."),
    ],
    n, FOOT,
)

# 26 ─ remediation
n += 1
table_slide(
    "Remediation order", "Security",
    ["#", "Action", "Why first"],
    [
        ["1", "Purge and ROTATE the committed secrets",
         "Untracking is not enough — the values are in history and must be treated as compromised"],
        ["2", "Generate SESSION_SECRET at bootstrap; refuse the placeholder",
         "Mechanical fix; closes cookie forgery and pending-token forgery together"],
        ["3", "Make OpenFGA authoritative in _require_role",
         "Removes the privilege-escalation half of the composed bypass"],
        ["4", "Fail closed when authorization is unavailable; add depends_on",
         "Turns a silent allow-all into a visible outage"],
        ["5", "Separate the OpenFGA management audience from the user audience",
         "Stops any user token from writing tuples"],
        ["6", "Terminate TLS at the portal; set SESSION_SECURE=true",
         "Prerequisite for exposing the portal beyond localhost"],
    ],
    n, "Items 2 and 3 should ship together — each alone leaves the composed bypass partly open.",
    widths=[0.05, 0.42, 0.53],
)

# 27 ─ closing
n += 1
closing_slide(
    "Where to read more",
    [
        ("Start here", "h"),
        ("ARCHITECTURE.md  —  repo root.  Topology, flows, authn/authz, full security posture.", "b"),
        ("identity_service/ARCHITECTURE.md  —  sessions, the OIDC exchange, enforcement point 1.", "b"),
        ("SCRIPTS.md  —  operator reference: bootstrap, verification, Vault, air-gap migration.", "b"),
        ("Per-subsystem detail", "h"),
        ("dex/ · lldap/ · vault/ · openfga_postgres/ · openfga_visualized/", "b"),
        ("libcloud.rest/ · libcloud/ · stoplight_mock/ · server/", "b"),
        ("A note on these documents", "h"),
        ("Every claim was verified against source. Where the existing docs disagreed with the code,", "b"),
        ("the code won and the documents were corrected — including a Vault LDAP auth method that was", "b"),
        ("documented in detail but does not exist, and a credential model described backwards.", "b"),
    ],
    n,
)

out_dir = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(out_dir, "libcloud_nutanix_system_overview.pptx")
prs.save(out)
print(f"wrote {out}  ({n} slides)")
